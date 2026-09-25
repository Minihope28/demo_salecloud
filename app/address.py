"""Adresse marocaine : ville, code postal et vérification.

Le référentiel (localité → code postal → province) doit être le fichier officiel de Barid Al-Maghrib,
publié sur le portail Open Data du Maroc (data.gov.ma, jeu « Codes postaux des localités »),
enregistré en CSV ou XLSX dans ``config/codes_postaux.csv`` (ou ``.xlsx``).
Les colonnes sont détectées d'après leur en-tête (code, localité/ville/commune, province, région).

Sans référentiel, l'application se contente de règles simples (code à 5 chiffres, ville déduite de la
fin de l'adresse) et le signale.
"""

from __future__ import annotations

import csv
import difflib
import io
import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

COUNTRY = "Maroc"
POSTAL_RE = re.compile(r"(?<!\d)(\d{5})(?!\d)")
STREET_WORDS = {"av", "ave", "avenue", "bd", "boulevard", "rue", "lot", "lotissement", "quartier", "qt", "hay",
                "residence", "res", "imm", "immeuble", "n", "no", "numero", "apt", "appt", "etage", "bloc", "zone",
                "route", "km", "douar", "angle", "cite", "place", "derb", "bp", "magasin", "local", "sidi", "ain"}


def norm(text: str) -> str:
    text = "".join(c for c in unicodedata.normalize("NFKD", text or "") if not unicodedata.combining(c)).lower()
    text = re.sub(r"[-'’_/]", " ", text)
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9 ]", " ", text)).strip()


@dataclass
class Locality:
    name: str
    codes: list[str] = field(default_factory=list)
    province: str = ""
    region: str = ""


class PostalReference:
    def __init__(self, localities: list[Locality], source: str):
        self.source = source
        self.by_name: dict[str, Locality] = {}
        self.by_code: dict[str, list[Locality]] = {}
        for loc in localities:
            key = norm(loc.name)
            if not key:
                continue
            existing = self.by_name.get(key)
            if existing:
                existing.codes += [c for c in loc.codes if c not in existing.codes]
                existing.province = existing.province or loc.province
                existing.region = existing.region or loc.region
            else:
                self.by_name[key] = loc
        for loc in self.by_name.values():
            for code in loc.codes:
                self.by_code.setdefault(code, []).append(loc)

    def __len__(self) -> int:
        return len(self.by_name)

    def get(self, city: str) -> Locality | None:
        return self.by_name.get(norm(city))

    def close_matches(self, city: str, n: int = 3) -> list[str]:
        keys = difflib.get_close_matches(norm(city), list(self.by_name), n=n, cutoff=0.75)
        return [self.by_name[k].name for k in keys]

    # ------------------------------------------------------------------ chargement
    @classmethod
    def load(cls, path: Path) -> "PostalReference":
        rows = _read_rows(path)
        if not rows:
            raise ValueError("fichier vide")
        header = [norm(str(h or "")) for h in rows[0]]

        def col(*needles: str) -> int | None:
            for i, h in enumerate(header):
                if any(n in h for n in needles):
                    return i
            return None

        code_i = col("code postal", "code")
        name_i = col("localite", "ville", "commune", "libelle", "nom")
        prov_i, reg_i = col("province", "prefecture"), col("region")
        if code_i is None or name_i is None:
            raise ValueError(f"colonnes « code » et « localité/ville » introuvables dans l'en-tête : {rows[0]}")

        localities = []
        for row in rows[1:]:
            if len(row) <= max(code_i, name_i):
                continue
            name = str(row[name_i] or "").strip()
            codes = POSTAL_RE.findall(str(row[code_i] or "").replace(".0", ""))
            if name and codes:
                localities.append(Locality(
                    name=name, codes=codes,
                    province=str(row[prov_i] or "").strip() if prov_i is not None and len(row) > prov_i else "",
                    region=str(row[reg_i] or "").strip() if reg_i is not None and len(row) > reg_i else "",
                ))
        if not localities:
            raise ValueError("aucune localité avec un code postal à 5 chiffres")
        return cls(localities, source=path.name)


def _read_rows(path: Path) -> list[list]:
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t")
    return list(csv.reader(io.StringIO(text), dialect))


def load_reference(paths: list[Path]) -> tuple[PostalReference | None, str]:
    """Charge le premier fichier existant. Retourne (référentiel, message d'état)."""
    for path in paths:
        if path.exists():
            try:
                ref = PostalReference.load(path)
            except (ValueError, OSError, csv.Error) as exc:
                return None, f"Référentiel des codes postaux illisible ({path.name}) : {exc}"
            return ref, f"Référentiel des codes postaux : {len(ref)} localités ({path.name})"
    return None, "Référentiel des codes postaux non chargé : seules des vérifications simples sont faites."


# =========================================================================== analyse

def _find_city(address: str, ref: PostalReference | None) -> tuple[str, str] | None:
    """Retourne (ville, explication) ou None."""
    text = norm(address)
    words = text.split()
    if ref:
        found: list[tuple[int, Locality, str]] = []
        for size in (4, 3, 2, 1):
            for i in range(len(words) - size + 1):
                chunk = " ".join(words[i:i + size])
                loc = ref.by_name.get(chunk)
                if loc and (len(chunk) >= 4 or size > 1):
                    found.append((i, loc, "exact"))
        # orthographe différente (ex. « TINJDAD » pour « Tinejdad ») : un mot long très proche
        for i, w in enumerate(words):
            if len(w) >= 5 and w not in ref.by_name and w not in STREET_WORDS:
                close = difflib.get_close_matches(w, list(ref.by_name), n=1, cutoff=0.85)
                if close:
                    found.append((i, ref.by_name[close[0]], "approchant"))
        if found:
            # Une localité située dans une province également citée est plus précise que la province.
            names = {norm(f[1].name) for f in found}
            specific = [f for f in found if norm(f[1].province) in names and norm(f[1].province) != norm(f[1].name)]
            pool = specific or found
            pos, loc, how = max(pool, key=lambda f: (f[2] == "exact" and not specific, f[0]))
            msg = "trouvée dans l'adresse" if how == "exact" else f"orthographe approchante dans l'adresse"
            return loc.name, f"{msg} (référentiel {ref.source})"

    # Sans référentiel : ville après le code postal, sinon dernier segment après une virgule.
    m = re.search(r"\d{5}\s+([A-Za-zÀ-ÿ' \-]{3,40})$", address.strip())
    if m:
        return m.group(1).strip().title(), "déduite de l'adresse (après le code postal)"
    last = address.rsplit(",", 1)[-1].strip() if "," in address else ""
    if last and re.fullmatch(r"[A-Za-zÀ-ÿ' \-]{3,40}", last) and norm(last).split()[0] not in STREET_WORDS:
        return last.title(), "déduite de l'adresse (dernier élément)"
    return None


def _strip_city(address: str, city: str, postal: str | None, province: str | None) -> str:
    """Adresse de rue : on retire en fin d'adresse la ville, le code postal et la province."""
    parts = [p.strip() for p in address.split(",")]
    drop = {norm(city), norm(province or "")} - {""}
    while len(parts) > 1 and (norm(POSTAL_RE.sub("", parts[-1])) in drop or norm(parts[-1]) == (postal or "")):
        parts.pop()
    street = ", ".join(parts)
    if postal:
        street = re.sub(rf"\s*{postal}\s*", " ", street)
    # ville collée en fin de rue : « AV HASSAN2 TINJDAD »
    words = street.split()
    if len(words) > 1:
        last = norm(words[-1])
        if last == norm(city) or difflib.SequenceMatcher(None, last, norm(city)).ratio() >= 0.85:
            words.pop()
    return " ".join(words).strip(" ,")


def analyse_address(address: str, ref: PostalReference | None) -> dict:
    """Propose rue, ville, code postal et province à partir d'une adresse complète."""
    result: dict = {"street": None, "city": None, "postal_code": None, "region": None, "notes": []}
    if not (address or "").strip():
        return result
    postal = (POSTAL_RE.findall(address) or [None])[-1]
    city = _find_city(address, ref)
    loc = ref.get(city[0]) if (ref and city) else None

    if city:
        result["city"] = {"value": city[0], "how": city[1]}
    if postal:
        result["postal_code"] = {"value": postal, "how": "trouvé dans l'adresse"}
    elif loc and len(loc.codes) == 1:
        result["postal_code"] = {"value": loc.codes[0], "how": f"déduit de la ville (référentiel {ref.source})"}
    elif loc and len(loc.codes) > 1:
        result["notes"].append(f"{loc.name} a plusieurs codes postaux ({', '.join(loc.codes[:6])}…) : à préciser selon le quartier.")
    if loc and (loc.province or loc.region):
        result["region"] = {"value": loc.region or loc.province, "how": f"référentiel {ref.source}"}
    if city:
        street = _strip_city(address, city[0], postal, loc.province if loc else None)
        if street and street != address.strip():
            result["street"] = {"value": street, "how": "adresse sans la ville ni le code postal"}
    return result


def check_postal(city: str, postal_code: str, ref: PostalReference | None) -> dict:
    """Vérifie la cohérence ville / code postal. status : ok | warning | error | unknown."""
    city, postal_code = (city or "").strip(), (postal_code or "").strip()
    if not postal_code or not city:
        return {"status": "unknown", "message": None, "suggestions": []}
    if not re.fullmatch(r"\d{5}", postal_code):
        return {"status": "error", "message": "Le code postal marocain comporte 5 chiffres.", "suggestions": []}
    if not ref:
        return {"status": "unknown", "message": "Référentiel des codes postaux non chargé : cohérence ville / code non vérifiée.",
                "suggestions": []}
    loc = ref.get(city)
    owners = [l.name for l in ref.by_code.get(postal_code, [])]
    if not loc:
        close = ref.close_matches(city)
        hint = f" Vouliez-vous dire : {', '.join(close)} ?" if close else ""
        return {"status": "warning", "message": f"Ville « {city} » absente du référentiel.{hint}", "suggestions": close}
    if postal_code in loc.codes:
        return {"status": "ok", "message": f"Code postal cohérent avec {loc.name}.", "suggestions": []}
    other = f" Le code {postal_code} correspond à : {', '.join(owners[:3])}." if owners else f" Le code {postal_code} est inconnu."
    return {"status": "warning",
            "message": f"Le code postal ne correspond pas à {loc.name} (attendu : {', '.join(loc.codes[:5])}).{other}",
            "suggestions": loc.codes[:5]}
